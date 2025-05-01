import openpyxl
from openpyxl import Workbook, load_workbook
import tkinter as tk
from tkinter import messagebox, ttk
import os

Fname = "Student_Score.xlsx"

def initial_workbook():
    if not os.path.exists(Fname):
        wb = Workbook()
        ws = wb.active
        ws.title = "S c o r e s"
        ws.append(["Name", "Score", "Status"])
        wb.save(Fname)
    
def get_status(score):
    return "Pass" if score >= 75 else "Fail"

def add_or_update(name, score):
    initial_workbook()
    wb = load_workbook(Fname)
    ws = wb.active

    update = False
    for row in ws.iter_row(min_row = 2, values_only = False):
        if row[0].value == name:
            row[1].value = score
            row[2].value = get_status(score)
            updated = True
            break

    if not updated:
        ws.append([name, score, get_status(score)])

    wb.save(Fname)
    return updated

def get_records():
    initial_workbook()
    wb = load_workbook(Fname)
    ws = wb.active

    records = []
    for row in ws.iter_rows(min_row = 2, values_only = True):
        records.append(row)
    return records

def submit_score():
    name = name_entry.get().strip
    try:
        score = int(score_entry.get().strip())
        if not (0 <= score <= 100):
            raise ValueError
    except ValueError:
        messagebox.showerror("Invalid Input!", "Score must be an integer between 0 and 100")
        return
    
    if name == "":
        messagebox.showerror("Invalid Input!", "Name cannot be empty!")
        return
    
    updated = add_or_update(name, score)
    message = f"Record {'updated' if updated else 'added'} for {name}."
    messagebox.showinfo("Success!", message)
    name_entry.delete(0, tk.END)
    score_entry.delete(0, tk.END)
    refresh_table()

def refresh_table():
    for row in tree.get_children():
        tree.delete(row)
    records = get_records()
    for record in records:
        tree.insert("", tk.END, values=record)

window = tk.Tk()
window.title("T w i s t y y ' s    S c o r e    T r a c k e r !")
window.configure(bg="Light Pink")
window.geometry("500x500")
window.resizable(False, False)

input_frame = tk.Frame(window)
input_frame.pack(pady=20)
input_frame.configure(bg="Pink")

tk.Label(input_frame, text="Student Name:").grid(row=0, column=0, padx=5, pady=5)
name_entry = tk.Entry(input_frame)
name_entry.grid(row=0, column=1, padx=5, pady=5)

tk.Label(input_frame, text="Score:").grid(row=1, column=0, padx=5, pady=5)
score_entry = tk.Entry(input_frame)
score_entry.grid(row=1, column=1, padx=5, pady=5)

submit_btn = tk.Button(input_frame, text="Submit Score", command=submit_score)
submit_btn.grid(row=2, column=0, columnspan=2, pady=10)


table_frame = tk.Frame(window)
table_frame.pack(pady=10)
table_frame.configure(bg="Pink")

tree = ttk.Treeview(table_frame, columns=("Name", "Score", "Status"), show="headings", height=10)
tree.heading("Name", text="Name")
tree.heading("Score", text="Score")
tree.heading("Status", text="Status")
tree.column("Name", width=150)
tree.column("Score", width=100)
tree.column("Status", width=100)
tree.pack()

refresh_table()

window.mainloop()